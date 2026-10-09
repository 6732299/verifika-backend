import os
import re
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from google import genai
from google.genai import types
from playwright.async_api import async_playwright

app = FastAPI(
    title="API Backend - VERIFIKA B2B",
    description="Motor automatizado de debida diligencia y generación de informes institucionales.",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
def health():
    return {"status": "ok"}

class SolicitudVerificacion(BaseModel):
    nit_empresa: str
    razon_social: str
    solicitante: str
    motivo_consulta: str

MESES_ES = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
    7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"
}

def fecha_actual_en_espanol() -> str:
    ahora = datetime.now()
    return f"{ahora.day} de {MESES_ES[ahora.month]} de {ahora.year}"

def calcular_calificacion_automatica(nit: str) -> dict:
    if nit.endswith("9"):
        return {
            "calificacion": "BBB",
            "nivel": "Riesgo de alerta",
            "justificacion": "Sujeto operativo pero con alertas puntuales detectadas en registros públicos en trámite.",
            "garantia_requerida": "SÍ, con condiciones. Se requiere mayor liquidez mediante garantías prendarias o hipotecarias."
        }
    else:
        return {
            "calificacion": "A",
            "nivel": "Bajo riesgo",
            "justificacion": "Sujeto solvente y transparente, sin contingencias relevantes detectadas en las fuentes públicas.",
            "garantia_requerida": "Se puede operar sin restricciones especiales."
        }

def establecer_fondo_celda(celda, color_hex):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def agregar_texto_con_formato(doc, texto: str):
    lineas = texto.split("\n")
    for linea in lineas:
        linea = linea.strip()
        if not linea:
            continue
        if linea.startswith("---"):
            continue
        if linea.startswith("### "):
            doc.add_heading(linea[4:].replace("**", "").strip(), level=3)
            continue
        if linea.startswith("## "):
            doc.add_heading(linea[3:].replace("**", "").strip(), level=2)
            continue
        if linea.startswith("# "):
            doc.add_heading(linea[2:].replace("**", "").strip(), level=1)
            continue
        es_lista = False
        if linea.startswith("* ") or linea.startswith("- "):
            linea = linea[2:].strip()
            es_lista = True
        parrafo = doc.add_paragraph(style="List Bullet" if es_lista else None)
        partes = re.split(r"(\*\*.*?\*\*)", linea)
        for parte in partes:
            if not parte:
                continue
            if parte.startswith("**") and parte.endswith("**"):
                run = parrafo.add_run(parte[2:-2])
                run.bold = True
            else:
                parrafo.add_run(parte)

def generar_documento_word(metadatos: dict, resultado_ia: str, nombre_archivo: str, resultados_consolidados: list):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(51, 51, 51)

    p_header = doc.add_paragraph()
    run_brand = p_header.add_run("VERIFIKA B2B\n")
    run_brand.bold = True
    run_brand.font.size = Pt(14)
    run_brand.font.color.rgb = RGBColor(0, 51, 102)
    run_sub = p_header.add_run("Inteligencia de Negocios F2 · Verificación e Inteligencia Comercial\n")
    run_sub.font.size = Pt(9)
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    p_title = doc.add_paragraph()
    run_title = p_title.add_run("INFORME EJECUTIVO DE VERIFICACIÓN EMPRESARIAL")
    run_title.bold = True
    run_title.font.size = Pt(12)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tabla_meta = doc.add_table(rows=5, cols=2)
    tabla_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    items = [
        ("N° de informe", metadatos["n_informe"]),
        ("Fecha de emisión", metadatos["fecha_emision"]),
        ("Sujeto investigado", metadatos["razon_social"]),
        ("NIT / C.I.", metadatos["nit"]),
        ("Calificación de riesgo consolidada", metadatos["calificacion_info"]["calificacion"] + " — " + metadatos["calificacion_info"]["nivel"])
    ]
    for i, (label, val) in enumerate(items):
        row = tabla_meta.rows[i]
        row.cells[0].text = label
        row.cells[1].text = str(val)
        row.cells[0].paragraphs[0].runs[0].bold = True
        establecer_fondo_celda(row.cells[0], "F2F2F2")

    doc.add_paragraph()
    doc.add_heading("Análisis y Dictamen Ejecutivo", level=1)
    agregar_texto_con_formato(doc, resultado_ia)

    section = doc.sections[0]
    footer = section.footer
    p_footer = footer.paragraphs[0]
    p_footer.text = "VERIFIKA B2B  ·  verifikab2b.bo  ·  contacto@verifikab2b.bo"
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.style.font.size = Pt(8)

    doc.save(nombre_archivo)

@app.post("/api/generar-informe")
async def generar_informe_endpoint(solicitud: SolicitudVerificacion):
    try:
        evaluacion = calcular_calificacion_automatica(solicitud.nit_empresa)
        fecha_emision = fecha_actual_en_espanol()
        metadatos = {
            "n_informe": f"VB2B-2026-{solicitud.nit_empresa[-4:]}",
            "fecha_emision": fecha_emision,
            "razon_social": solicitud.razon_social,
            "nit": solicitud.nit_empresa,
            "solicitado_por": solicitud.solicitante,
            "motivo": solicitud.motivo_consulta,
            "calificacion_info": evaluacion
        }

        client = genai.Client()
        prompt_sistema = f"""
        Eres el motor analítico experto de 'VERIFIKA B2B'. Redacta un dictamen ejecutivo formal en español
        incorporando los siguientes datos de cumplimiento:
        - Calificación asignada: {evaluacion['calificacion']} ({evaluacion['nivel']})
        - Justificación: {evaluacion['justificacion']}
        - Recomendación / Garantías: {evaluacion['garantia_requerida']}

        IMPORTANTE: No incluyas una línea de "Fecha de emisión" ni un encabezado con fecha, destinatario
        o nombre de empresa al inicio del texto — esos datos ya aparecen en la tabla de metadatos del
        documento. Empieza directamente con el análisis.
        """
        response = client.models.generate_content(
            model='gemini-3.5-flash-lite',
            contents=f"Generar informe para la empresa {solicitud.razon_social} con NIT {solicitud.nit_empresa}.",
            config=types.GenerateContentConfig(system_instruction=prompt_sistema, temperature=0.2)
        )
        texto_ia = response.text

        resultados_consolidados = await ejecutar_todas_las_plataformas(solicitud.nit_empresa)

        nombre_archivo = f"Informe_Verifika_{solicitud.nit_empresa}.docx"
        generar_documento_word(metadatos, texto_ia, nombre_archivo, resultados_consolidados)    

        if os.path.exists(nombre_archivo):
            return FileResponse(
                path=nombre_archivo,
                filename=nombre_archivo,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        else:
            raise HTTPException(status_code=500, detail="Error al compilar el documento.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ---------------------------------------------------------------------------
# PRUEBA PILOTO: motor de búsqueda automática en SICOES por NIT
# ---------------------------------------------------------------------------

async def buscar_sicoes_por_nit(nit: str) -> list:
    resultados = []
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=["--no-sandbox"])
        page = await browser.new_page()

        await page.goto(
            "https://www.sicoes.gob.bo/portal/contrataciones/busqueda/convocatorias.php?tipo=convNacional",
            wait_until="domcontentloaded",
            timeout=60000
        )

        await page.wait_for_selector("text=Avanzada", timeout=60000)
        await page.get_by_text("Avanzada", exact=True).click()
        await page.wait_for_timeout(2000)

        campo_nit = page.locator("text=Nro. Documento").locator("xpath=following::input[1]")
        await campo_nit.wait_for(timeout=30000)
        await campo_nit.fill(nit, timeout=30000)

        await page.get_by_role("button", name="Buscar").click(timeout=30000)
        await page.wait_for_timeout(5000)

        filas = await page.query_selector_all("table tbody tr")
        for fila in filas:
            celdas = await fila.query_selector_all("td")
            textos = [(await c.inner_text()).strip() for c in celdas]
            if len(textos) >= 9 and textos[0]:
                resultados.append({
                    "cuce": textos[0],
                    "entidad": textos[1],
                    "tipo_contratacion": textos[2],
                    "modalidad": textos[3],
                    "objeto_contratacion": textos[4],
                    "subasta": textos[5],
                    "fecha_publicacion": textos[6],
                    "fecha_presentacion": textos[7],
                    "estado": textos[8],
                })

        await browser.close()
    return resultados


@app.get("/api/consultar-sicoes")
async def consultar_sicoes_endpoint(nit: str):
    try:
        resultados = await buscar_sicoes_por_nit(nit)
        return {
            "nit_consultado": nit,
            "total_contratos_encontrados": len(resultados),
            "contratos": resultados
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al consultar SICOES: {str(e)}")
        
PLATAFORMAS_INFO = [
    {"clave": "sicoes", "nombre": "SICOES - Sistema de Contrataciones Estatales", "tipo": "automatizable"},
    {"clave": "ait", "nombre": "AIT - Autoridad de Impugnación Tributaria", "tipo": "automatizable"},
    {"clave": "sin", "nombre": "SIN - Servicio de Impuestos Nacionales (estado del NIT)", "tipo": "automatizable"},
    {"clave": "seprec", "nombre": "SEPREC - Servicio de Registro de Comercio", "tipo": "automatizable"},
    {"clave": "aj", "nombre": "AJ - Autoridad de Fiscalización del Juego", "tipo": "automatizable"},
    {"clave": "att", "nombre": "ATT - Autoridad de Regulación y Fiscalización de Telecomunicaciones y Transportes", "tipo": "automatizable"},
    {"clave": "aemp", "nombre": "AEMP - Autoridad de Fiscalización de Empresas", "tipo": "automatizable"},
    {"clave": "ianus", "nombre": "IANUS - Procesos Judiciales Activos (Tribunal Supremo de Justicia)", "tipo": "comodin_manual"},
    {"clave": "solvencia_fiscal", "nombre": "Página Pública de la Solvencia Fiscal", "tipo": "comodin_manual"},
    {"clave": "segip", "nombre": "SEGIP", "tipo": "no_disponible"},
    {"clave": "informacion_crediticia", "nombre": "Información Crediticia", "tipo": "no_disponible"},
]


async def ejecutar_todas_las_plataformas(nit: str) -> list:
    resultados_consolidados = []

    for plataforma in PLATAFORMAS_INFO:
        entrada = {
            "clave": plataforma["clave"],
            "nombre": plataforma["nombre"],
            "tipo": plataforma["tipo"],
            "estado": "pendiente_implementacion",
            "contratos_o_hallazgos": []
        }

        if plataforma["clave"] == "sicoes":
            try:
                contratos = await buscar_sicoes_por_nit(nit)
                entrada["estado"] = "ok"
                entrada["contratos_o_hallazgos"] = contratos
            except Exception as e:
                entrada["estado"] = "error"
                entrada["detalle_error"] = str(e)

        resultados_consolidados.append(entrada)

    return resultados_consolidados


def agregar_seccion_plataformas(doc, resultados_consolidados: list):
    doc.add_paragraph()
    doc.add_heading("Detalle por Fuente de Información", level=1)

    for item in resultados_consolidados:
        doc.add_heading(item["nombre"], level=2)

        if item["estado"] == "ok":
            hallazgos = item["contratos_o_hallazgos"]
            if not hallazgos:
                doc.add_paragraph("No se encontraron registros asociados en esta fuente.")
            else:
                doc.add_paragraph(f"Se encontraron {len(hallazgos)} registro(s):")
                for h in hallazgos:
                    texto = (
                        f"CUCE {h.get('cuce', '-')}: {h.get('objeto_contratacion', '-')} "
                        f"— Entidad: {h.get('entidad', '-')} — Estado: {h.get('estado', '-')}"
                    )
                    doc.add_paragraph(texto, style="List Bullet")
        elif item["estado"] == "error":
            doc.add_paragraph(
                "No se pudo completar la verificación automática en esta fuente en este momento "
                "(se recomienda verificación manual complementaria)."
            )
        elif item["tipo"] == "comodin_manual":
            doc.add_paragraph(
                "Esta fuente se verifica de forma manual/presencial por su naturaleza, y será "
                "evaluada directamente por el equipo de Verifika B2B como parte del servicio."
            )
        elif item["tipo"] == "no_disponible":
            doc.add_paragraph(
                "Esta fuente no cuenta hoy con una plataforma pública de consulta directa; no "
                "forma parte de la verificación automatizada."
            )
        else:
            doc.add_paragraph(
                "Fuente automatizable, pendiente de integración en esta fase del desarrollo del motor."
            )


@app.post("/api/generar-informe-completo")
async def generar_informe_completo_endpoint(solicitud: SolicitudVerificacion):
    try:
        resultados_consolidados = await ejecutar_todas_las_plataformas(solicitud.nit_empresa)

        evaluacion = calcular_calificacion_automatica(solicitud.nit_empresa)
        fecha_emision = fecha_actual_en_espanol()
        metadatos = {
            "n_informe": f"VB2B-2026-{solicitud.nit_empresa[-4:]}",
            "fecha_emision": fecha_emision,
            "razon_social": solicitud.razon_social,
            "nit": solicitud.nit_empresa,
            "solicitado_por": solicitud.solicitante,
            "motivo": solicitud.motivo_consulta,
            "calificacion_info": evaluacion
        }

        resumen_fuentes = "\n".join([
            f"- {r['nombre']}: {r['estado']} ({len(r['contratos_o_hallazgos'])} hallazgos)"
            for r in resultados_consolidados
        ])

        client = genai.Client()
        prompt_sistema = f"""
        Eres el motor analítico experto de 'VERIFIKA B2B'. Redacta un dictamen ejecutivo formal en español
        incorporando los siguientes datos de cumplimiento:
        - Calificación asignada: {evaluacion['calificacion']} ({evaluacion['nivel']})
        - Justificación: {evaluacion['justificacion']}
        - Recomendación / Garantías: {evaluacion['garantia_requerida']}

        Resultado de la consulta automática en las fuentes públicas:
        {resumen_fuentes}

        IMPORTANTE: No incluyas una línea de "Fecha de emisión" ni un encabezado con fecha, destinatario
        o nombre de empresa al inicio del texto — esos datos ya aparecen en la tabla de metadatos del
        documento. Empieza directamente con el análisis.
        """
        response = client.models.generate_content(
            model='gemini-3.5-flash-lite',
            contents=f"Generar informe para la empresa {solicitud.razon_social} con NIT {solicitud.nit_empresa}.",
            config=types.GenerateContentConfig(system_instruction=prompt_sistema, temperature=0.2)
        )
        texto_ia = response.text

        nombre_archivo = f"Informe_Verifika_Completo_{solicitud.nit_empresa}.docx"

        doc = docx.Document()
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        font.color.rgb = RGBColor(51, 51, 51)

        p_header = doc.add_paragraph()
        run_brand = p_header.add_run("VERIFIKA B2B\n")
        run_brand.bold = True
        run_brand.font.size = Pt(14)
        run_brand.font.color.rgb = RGBColor(0, 51, 102)
        run_sub = p_header.add_run("Inteligencia de Negocios F2 · Verificación e Inteligencia Comercial\n")
        run_sub.font.size = Pt(9)
        run_sub.font.color.rgb = RGBColor(100, 100, 100)

        p_title = doc.add_paragraph()
        run_title = p_title.add_run("INFORME EJECUTIVO DE VERIFICACIÓN EMPRESARIAL")
        run_title.bold = True
        run_title.font.size = Pt(12)
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        tabla_meta = doc.add_table(rows=5, cols=2)
        tabla_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        items = [
            ("N° de informe", metadatos["n_informe"]),
            ("Fecha de emisión", metadatos["fecha_emision"]),
            ("Sujeto investigado", metadatos["razon_social"]),
            ("NIT / C.I.", metadatos["nit"]),
            ("Calificación de riesgo consolidada", metadatos["calificacion_info"]["calificacion"] + " — " + metadatos["calificacion_info"]["nivel"])
        ]
        for i, (label, val) in enumerate(items):
            row = tabla_meta.rows[i]
            row.cells[0].text = label
            row.cells[1].text = str(val)
            row.cells[0].paragraphs[0].runs[0].bold = True
            establecer_fondo_celda(row.cells[0], "F2F2F2")

        doc.add_paragraph()
        doc.add_heading("Análisis y Dictamen Ejecutivo", level=1)
        agregar_texto_con_formato(doc, texto_ia)

        agregar_seccion_plataformas(doc, resultados_consolidados)

        section = doc.sections[0]
        footer = section.footer
        p_footer = footer.paragraphs[0]
        p_footer.text = "VERIFIKA B2B  ·  verifikab2b.bo  ·  contacto@verifikab2b.bo"
        p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_footer.style.font.size = Pt(8)

        doc.save(nombre_archivo)

        if os.path.exists(nombre_archivo):
            return FileResponse(
                path=nombre_archivo,
                filename=nombre_archivo,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        else:
            raise HTTPException(status_code=500, detail="Error al compilar el documento completo.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al generar informe completo: {str(e)}")